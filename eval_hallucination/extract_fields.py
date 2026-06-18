# -*- coding: utf-8 -*-
"""
Step 1: Extract fields to be judged.

For each (raw, parser) pair, produce a JSON file listing every atomic string
value in the parsed JSON, paired with the full source text. The LLM judge
(me) will later emit a verdict for every value.

Output: eval_hallucination/fields_for_judge.json
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path
from typing import Any, Iterable

# ── text extraction backends ────────────────────────────────────────────────
def _extract_pdf(path: Path) -> str:
    try:
        import pdfplumber
        with pdfplumber.open(str(path)) as pdf:
            return "\n".join((p.extract_text() or "") for p in pdf.pages)
    except Exception as e:
        print(f"[warn] pdfplumber failed on {path.name}: {e}", file=sys.stderr)
        return ""

def _extract_docx(path: Path) -> str:
    try:
        import docx
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for tbl in d.tables:
            for row in tbl.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        return "\n".join(parts)
    except Exception as e:
        print(f"[warn] docx failed on {path.name}: {e}", file=sys.stderr)
        return ""

def _extract_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")

_EXTRACTORS = {".pdf": _extract_pdf, ".docx": _extract_docx, ".txt": _extract_txt}

def extract_text(path: Path) -> str:
    fn = _EXTRACTORS.get(path.suffix.lower())
    if not fn:
        return ""
    return fn(path)

# ── filename parsing ──────────────────────────────────────────────────────
_NAME_RE = re.compile(r"^(cv|jd|ci)_(\d+)_(\d+)_(\d+)\.(.+)$", re.IGNORECASE)

def parse_name(name: str) -> tuple[str, int, int, str] | None:
    m = _NAME_RE.match(name)
    if not m:
        return None
    ftype, _ts, user_id, rec_id, ext = m.groups()
    return ftype.lower(), int(user_id), int(rec_id), ext.lower()

# ── walk JSON, collect atomic values with their path ──────────────────────
def walk_strings(obj: Any, path: str = "") -> Iterable[tuple[str, str]]:
    if isinstance(obj, dict):
        for k, v in obj.items():
            yield from walk_strings(v, f"{path}.{k}" if path else str(k))
    elif isinstance(obj, list):
        for i, v in enumerate(obj):
            yield from walk_strings(v, f"{path}[{i}]")
    elif isinstance(obj, str):
        s = obj.strip()
        if s:
            yield path, s

# ── main ──────────────────────────────────────────────────────────────────
def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument(
        "--raw-dir",
        type=Path,
        default=Path(__file__).resolve().parent.parent
        / "app/feature/feature_up_cv/storage/raw_file",
    )
    ap.add_argument(
        "--parser-dir",
        type=Path,
        default=Path(__file__).resolve().parent.parent
        / "app/feature/feature_up_cv/storage/parser_file",
    )
    ap.add_argument(
        "--out",
        type=Path,
        default=Path(__file__).resolve().parent / "fields_for_judge.json",
    )
    ap.add_argument(
        "--max-chars",
        type=int,
        default=20000,
        help="Truncate source text per doc to this many characters "
             "(LLM context budget; raise if you have more budget).",
    )
    args = ap.parse_args()

    if not args.raw_dir.is_dir():
        print(f"[fatal] raw-dir not found: {args.raw_dir}", file=sys.stderr)
        return 2
    if not args.parser_dir.is_dir():
        print(f"[fatal] parser-dir not found: {args.parser_dir}", file=sys.stderr)
        return 2

    parser_by_key: dict[tuple[str, int, int], Path] = {}
    for f in args.parser_dir.iterdir():
        if f.suffix.lower() != ".json":
            continue
        meta = parse_name(f.name)
        if not meta:
            continue
        ftype, uid, rid, _ = meta
        parser_by_key[(ftype, uid, rid)] = f

    documents: list[dict[str, Any]] = []
    for raw in sorted(args.raw_dir.iterdir()):
        if not raw.is_file():
            continue
        meta = parse_name(raw.name)
        if not meta:
            continue
        ftype, uid, rid, _ = meta
        parser = parser_by_key.get((ftype, uid, rid))
        if not parser:
            print(f"[skip] no parser file for raw {raw.name}", file=sys.stderr)
            continue

        source_text = extract_text(raw)
        if not source_text:
            print(f"[warn] empty text for {raw.name}", file=sys.stderr)
            continue

        try:
            parsed = json.loads(parser.read_text(encoding="utf-8"))
        except Exception as e:
            print(f"[warn] cannot parse {parser.name}: {e}", file=sys.stderr)
            continue

        fields = [
            {"id": f"{raw.stem}::{path}", "path": path, "value": value}
            for path, value in walk_strings(parsed)
        ]

        source_truncated = source_text[: args.max_chars]
        documents.append(
            {
                "doc_id": raw.stem,
                "raw_file": raw.name,
                "parser_file": parser.name,
                "file_type": ftype,
                "user_id": uid,
                "record_id": rid,
                "source_chars": len(source_text),
                "source_chars_truncated_to": len(source_truncated),
                "source_text": source_truncated,
                "fields": fields,
            }
        )

    args.out.write_text(
        json.dumps(documents, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    total_fields = sum(len(d["fields"]) for d in documents)
    print(
        f"[ok] wrote {len(documents)} docs, {total_fields} fields to {args.out}",
        file=sys.stderr,
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
